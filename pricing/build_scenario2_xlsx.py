"""Build pricing/Asia-Digital-Bridge-Pricer.xlsx (Scenario 2) with live formulas.
Monte Carlo outputs are read from pricing/scenario2_results.json (run scenario2_vt_note.py first).
The file has no cached values; Excel recalculates on open (fullCalcOnLoad).
"""
import json
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

R = json.load(open("pricing/scenario2_results.json"))
OUT = "pricing/Asia-Digital-Bridge-Pricer.xlsx"
F = "Arial"
BLUE, GREEN, BLACK = "0000FF", "008000", "000000"
NAVY = "123A5A"
f_in = Font(name=F, color=BLUE); f_calc = Font(name=F, color=BLACK); f_link = Font(name=F, color=GREEN)
f_b = Font(name=F, bold=True); f_h = Font(name=F, bold=True, color="FFFFFF"); f_t = Font(name=F, bold=True, size=14, color=NAVY)
f_note = Font(name=F, italic=True, size=9, color="666666")
fill_h = PatternFill("solid", fgColor=NAVY); fill_y = PatternFill("solid", fgColor="FFFF00")
fill_k = PatternFill("solid", fgColor="E8F0F7")
thin = Border(bottom=Side(style="thin", color="BBBBBB"))
PCT, PCT0, NUM4, USD, MULT = "0.00%", "0%", "0.0000", '$#,##0;($#,##0);-', "0.00x"

wb = Workbook()


def sheet(name, title, widths):
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = f_t
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    return ws


def header(ws, row, labels, col=1):
    for j, l in enumerate(labels):
        c = ws.cell(row, col + j, l); c.font = f_h; c.fill = fill_h; c.alignment = Alignment(horizontal="center", wrap_text=True)


def put(ws, ref, v, font=f_calc, fmt=None, fill=None, note=None):
    c = ws[ref]; c.value = v; c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if note: c.comment = Comment(note, "Team")
    return c


# ---------------- Pricer ----------------
ws = wb.active; ws.title = "Pricer"
ws["A1"] = "Asia Digital Bridge Note: 7Y USD 100% capital-protected note on a 10% vol-target index"; ws["A1"].font = f_t
for i, w in enumerate([44, 16, 4, 12, 12, 4, 60], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.sheet_view.showGridLines = False
put(ws, "A2", "Blue = input (edit these), black = formula, green = link to another sheet. Yellow = key assumption to verify on Bloomberg as of 17 Sep 2026.", f_note)

header(ws, 4, ["Inputs", "Value"])
inputs = [
    ("Notional (USD)", 100_000_000, USD, "Client mandate: USD 100Mn (rules, Scenario 2)", False),
    ("Tenor (years)", 7, "0", "Team choice: 7Y = one-generation handover window", False),
    ("Capital protection at maturity", 1.0, PCT0, "100% of notional, subject to Natixis credit", False),
    ("Issuer fee / margin (upfront)", 0.02, PCT, "Assumption: 2% over the life (~29bp p.a.)", True),
    ("USD rate for option pricing (OIS)", 0.0375, PCT, "Assumption: USD OIS/SOFR 7Y; replace with Bloomberg USSO7 Curncy", True),
    ("VT index: target volatility", 0.10, PCT, "Index design choice (see Design Dial)", False),
    ("VT index: decrement p.a.", 0.01, PCT, "Index design choice (see Design Dial)", False),
    ("MC / Black-Scholes price ratio", R["c_mc"] / R["c_bs"], NUM4,
     "From scenario2_vt_note.py: Monte Carlo price of the call (regime-switching basket, vol-target lag) / BS price at target vol", True),
]
for i, (lab, v, fmt, note, key) in enumerate(inputs):
    r = 5 + i
    put(ws, f"A{r}", lab, f_b if False else Font(name=F))
    put(ws, f"B{r}", v, f_in, fmt, fill_y if key else None, note)
# names for readability
N_, T_, P_, FEE_, R_, V_, D_, ADJ_ = (f"$B${r}" for r in range(5, 13))

header(ws, 4, ["Tenor", "USD funding"], col=4)
grid = [(1, .0478), (2, .053), (3, .0552), (5, .0569), (7, .0575), (10, .0582), (20, .0604)]
for i, (t, f) in enumerate(grid):
    put(ws, f"D{5 + i}", t, f_in, "0"); put(ws, f"E{5 + i}", f, f_in, PCT)
put(ws, "D12", "Source: 2026 game rules, Natixis funding grid (USD).", f_note)
TT, FF = "$D$5:$D$11", "$E$5:$E$11"

header(ws, 14, ["Pricing", "Value"])
rows = [
    ("Interpolated USD funding rate",
     f"=IF({T_}>=MAX({TT}),INDEX({FF},ROWS({TT})),INDEX({FF},MATCH({T_},{TT},1))+(INDEX({FF},MATCH({T_},{TT},1)+1)-INDEX({FF},MATCH({T_},{TT},1)))*({T_}-INDEX({TT},MATCH({T_},{TT},1)))/(INDEX({TT},MATCH({T_},{TT},1)+1)-INDEX({TT},MATCH({T_},{TT},1))))",
     PCT, "Linear interpolation on the grid (rules, footnote 1)"),
    ("Zero-coupon bond price (% notional)", f"={P_}/(1+B15)^{T_}", PCT, "Protection / (1 + funding)^T"),
    ("Option budget (% notional)", f"=1-B16-{FEE_}", PCT, "100% - bond - fee"),
    ("BS d1", f"=(({R_}-{D_})+0.5*{V_}^2)*{T_}/({V_}*SQRT({T_}))", NUM4, "VT index forward drifts at r - decrement"),
    ("BS d2", f"=B18-{V_}*SQRT({T_})", NUM4, ""),
    ("ATM call on VT index, Black-Scholes (% notional)", f"=EXP(-{D_}*{T_})*NORMSDIST(B18)-EXP(-{R_}*{T_})*NORMSDIST(B19)", PCT, ""),
    ("ATM call on VT index, MC-adjusted", f"=B20*{ADJ_}", PCT, "Captures vol-target lag / gap risk"),
    ("Participation (Black-Scholes)", f"=IF(B20>0,B17/B20,0)", PCT0, ""),
    ("PARTICIPATION (MC-adjusted, quoted)", f"=IF(B21>0,B17/B21,0)", PCT0, "Headline term for the term sheet"),
    ("Bond leg (USD)", f"={N_}*B16", USD, ""),
    ("Option leg (USD)", f"={N_}*B17", USD, ""),
    ("Issuer fee (USD)", f"={N_}*{FEE_}", USD, ""),
    ("Check: legs sum to notional", f"=IF(ABS(B24+B25+B26-{N_})<1,\"OK\",\"CHECK\")", None, ""),
]
for i, (lab, fml, fmt, note) in enumerate(rows):
    r = 15 + i
    put(ws, f"A{r}", lab, f_b if "PARTICIPATION" in lab else Font(name=F))
    c = put(ws, f"B{r}", fml, f_b if "PARTICIPATION" in lab else f_calc, fmt, fill_k if "PARTICIPATION" in lab else None)
    put(ws, f"G{r}", note, f_note)

# ---------------- Payoff ----------------
wp = sheet("Payoff", "Redemption at maturity vs VT index performance", [22, 18, 18, 18, 18])
put(wp, "A2", "Participation and protection link to the Pricer sheet. Direct basket shown for a VT-index move only (illustrative).", f_note)
put(wp, "A3", "Participation"); put(wp, "B3", "=Pricer!B23", f_link, PCT0)
put(wp, "C3", "Protection"); put(wp, "D3", "=Pricer!B7", f_link, PCT0)
put(wp, "A4", "Tenor (years)"); put(wp, "B4", "=Pricer!B6", f_link, "0")
header(wp, 6, ["VT index performance", "Redemption (% notional)", "Annualised return", "Payout (USD)", "Gain vs notional (USD)"])
perfs = [-0.4, -0.2, -0.1, 0, 0.1, 0.2, 0.3, 0.5, 0.75, 1.0, 1.5]
for i, p in enumerate(perfs):
    r = 7 + i
    put(wp, f"A{r}", p, f_in, PCT0)
    put(wp, f"B{r}", f"=$D$3+$B$3*MAX(0,A{r})", fmt=PCT)
    put(wp, f"C{r}", f"=B{r}^(1/$B$4)-1", fmt=PCT)
    put(wp, f"D{r}", f"=B{r}*Pricer!$B$5", fmt=USD)
    put(wp, f"E{r}", f"=D{r}-Pricer!$B$5", fmt=USD)

# ---------------- Design Dial ----------------
wd = sheet("Design Dial", "Participation by index design (Black-Scholes x MC ratio; links to Pricer)", [20, 13, 13, 13, 13, 13])
put(wd, "A2", "Higher decrement lowers the index forward and buys a bigger headline participation, but adds drag. See MC Results for what the client actually receives.", f_note)
header(wd, 4, ["Target vol \\ decrement", 0, 0.01, 0.02, 0.03, 0.04])
for j in range(5):
    wd.cell(4, 2 + j).number_format = PCT0
for i, tv in enumerate([0.08, 0.10, 0.12, 0.15]):
    r = 5 + i
    put(wd, f"A{r}", tv, f_in, PCT0)
    for j in range(5):
        col = get_column_letter(2 + j)
        d, v, t, rr = f"{col}$4", f"$A{r}", "Pricer!$B$6", "Pricer!$B$9"
        d1 = f"(({rr}-{d})+0.5*{v}^2)*{t}/({v}*SQRT({t}))"
        call = f"(EXP(-{d}*{t})*NORMSDIST({d1})-EXP(-{rr}*{t})*NORMSDIST({d1}-{v}*SQRT({t})))*Pricer!$B$12"
        put(wd, f"{col}{r}", f"=Pricer!$B$17/{call}", fmt=PCT0)

# ---------------- MC Results ----------------
wm = sheet("MC Results", "Monte Carlo outputs (values pasted from pricing/scenario2_vt_note.py)", [36, 13, 13, 13, 13, 13, 13, 13])
put(wm, "A2", "Placeholder market inputs (regime vol 27%/50%, OIS 3.75%, raw-basket implied vol 33%). Forward simulation, NOT a historical backtest.", f_note)
put(wm, "A4", "Feature menu (same simulated paths)", f_b)
header(wm, 5, ["Feature", "Option cost", "Participation"])
for i, (n, c, p) in enumerate(R["menu"]):
    put(wm, f"A{6 + i}", n); put(wm, f"B{6 + i}", c, f_in, PCT); put(wm, f"C{6 + i}", p, f_in, PCT0)
put(wm, "A11", "The decrement trap: headline vs outcome", f_b)
header(wm, 12, ["Index design", "Participation", "Median @6%", "P(only 100 back) @6%", "Median @9%", "P(only 100 back) @9%"])
for i, row in enumerate(R["designs"]):
    r = 13 + i
    put(wm, f"A{r}", row[0]); put(wm, f"B{r}", row[3], f_in, PCT0)
    for j, (v, fmt) in enumerate(zip(row[4:], [MULT, PCT0, MULT, PCT0])):
        put(wm, f"{get_column_letter(3 + j)}{r}", v, f_in, fmt)
put(wm, "A19", "Real-world outcomes at 7Y (multiple of capital)", f_b)
header(wm, 20, ["Basket return / product", "Mean", "5th pct", "Median", "95th pct", "Median CAGR", "P(below 100)", "P(100 or less)"])
r = 21
for mu, res in R["realworld"].items():
    for k, s in res.items():
        put(wm, f"A{r}", f"{float(mu):.0%} p.a. | {k}")
        for j, (key, fmt) in enumerate([("mean", MULT), ("p5", MULT), ("p50", MULT), ("p95", MULT), ("cagr50", PCT), ("p_loss", PCT0), ("p_floor", PCT0)]):
            put(wm, f"{get_column_letter(2 + j)}{r}", s[key], f_in, fmt)
        r += 1
put(wm, f"A{r + 1}", f"7Y UST proxy multiple: {R['ust_multiple']:.2f}x. Mean VT exposure {R['w_mean']:.0%}; realised index vol {R['idx_vol']:.1%}.", f_note)

# ---------------- Risk ----------------
wr = sheet("Risk", "Unit 15 risk section: VaR/ES, stress, sizing", [62, 14, 14, 14, 16, 14])
k = R["risk"]
put(wr, "A3", "Note fair value at issue (% notional)"); put(wr, "B3", k["v0"], f_in, PCT)
put(wr, "A4", "Delta: note change per 1% move in VT index"); put(wr, "B4", k["delta"], f_in, PCT)
put(wr, "A5", "DV01: note change per +1bp rates"); put(wr, "B5", k["dv01"], f_in, "0.000%")
put(wr, "A6", "10-day 99% VaR (delta-normal, % notional)"); put(wr, "B6", k["var99"], f_in, PCT)
put(wr, "A7", "10-day 99% Expected Shortfall (% notional)"); put(wr, "B7", k["es99"], f_in, PCT)
put(wr, "A8", "10-day 99% VaR (USD)"); put(wr, "B8", "=B6*Pricer!$B$5", fmt=USD)
header(wr, 10, ["Stress (spot, vol and rates shocked together)", "VT index level", "Rate shift", "Note MTM", "MTM P&L (USD)", "% of net worth"])
put(wr, "A9", "Family net worth (USD)"); put(wr, "B9", k["sizing"]["nw"], f_in, USD, note="Rules: combined net worth exceeding USD 2Bn")
for i, (n, s, dr, v) in enumerate(k["stress"]):
    r = 11 + i
    put(wr, f"A{r}", n); put(wr, f"B{r}", s, f_in, PCT); put(wr, f"C{r}", dr, f_in, "0.00%")
    put(wr, f"D{r}", v, f_in, PCT); put(wr, f"E{r}", f"=(D{r}-$B$3)*Pricer!$B$5", fmt=USD)
    put(wr, f"F{r}", f"=E{r}/$B$9", fmt=PCT)
put(wr, "A16", "Sizing identity (Unit 15)", f_b)
put(wr, "A17", "Allocation as % of net worth"); put(wr, "B17", "=Pricer!B5/B9", fmt=PCT)
put(wr, "A18", "Contractual max loss at maturity (ex issuer default)"); put(wr, "B18", "=1-Pricer!B7", fmt=PCT)
put(wr, "A19", "Worst stress MTM loss (USD) = notional x (fair value - worst MTM)"); put(wr, "B19", "=-MIN(E11:E14)", fmt=USD)
put(wr, "A20", "Worst stress MTM loss as % of net worth"); put(wr, "B20", "=B19/B9", fmt=PCT)
put(wr, "A22", "Exits: (1) profit target: sell back on the issuer bid once the index is up 60%; (2) no stop-loss on a protected note, the floor is the stop; "
              "(3) thesis invalidation: Taiwan blockade or sweeping chip export bans; (4) time: hold to maturity.", f_note)

# ---------------- Basket ----------------
wbk = sheet("Basket", "Underlying basket: Asian AI-hardware leaders (equal weight, USD quanto)", [14, 22, 10, 22, 12, 14, 14, 12])
put(wbk, "A2", "Fill the yellow columns from Bloomberg (CUR_MKT_CAP in USD Mn; 6M average daily traded value in USD Mn). Example values in row 5 are illustrative only.", f_note)
header(wbk, 4, ["Bloomberg", "Name", "Country", "Theme", "Weight", "Mkt cap USD Mn", "6M ADV USD Mn", "Eligible?"])
names = [("2330 TT", "TSMC", "TW", "Foundry / AI accelerators"), ("2454 TT", "MediaTek", "TW", "Custom AI ASICs / edge AI"),
         ("000660 KS", "SK Hynix", "KR", "HBM memory"), ("005930 KS", "Samsung Electronics", "KR", "HBM / memory"),
         ("8035 JT", "Tokyo Electron", "JP", "Wafer-fab equipment"), ("6857 JT", "Advantest", "JP", "AI chip testing"),
         ("6146 JT", "Disco", "JP", "Dicing / grinding for HBM"), ("4063 JT", "Shin-Etsu Chemical", "JP", "Silicon wafers / materials"),
         ("3711 TT", "ASE Technology", "TW", "Advanced packaging"), ("4062 JT", "Ibiden", "JP", "AI chip substrates"),
         ("2317 TT", "Hon Hai (Foxconn)", "TW", "AI server assembly"), ("2382 TT", "Quanta Computer", "TW", "AI servers"),
         ("2308 TT", "Delta Electronics", "TW", "Data-centre power / cooling")]
for i, (t, n, c, th) in enumerate(names):
    r = 5 + i
    put(wbk, f"A{r}", t, f_in); put(wbk, f"B{r}", n, f_in); put(wbk, f"C{r}", c, f_in); put(wbk, f"D{r}", th, f_in)
    put(wbk, f"E{r}", f"=1/COUNTA($A$5:$A${4 + len(names)})", fmt="0.0%")
    put(wbk, f"F{r}", 950000 if i == 0 else None, f_in, "#,##0", fill_y)
    put(wbk, f"G{r}", 4000 if i == 0 else None, f_in, "#,##0", fill_y)
    put(wbk, f"H{r}", f'=IF(OR(F{r}="",G{r}=""),"fill",IF(AND(F{r}>=250,G{r}>=5),"OK","FAIL"))')
end = 4 + len(names)
put(wbk, f"A{end + 1}", "Total weight"); put(wbk, f"E{end + 1}", f"=SUM(E5:E{end})", fmt="0.0%")
put(wbk, f"A{end + 3}", "Excluded on purpose: data-centre operators/REITs, telecom towers, renewables, real estate (overlap with family businesses); names on US restriction lists (e.g. SMIC). "
                        "Rules: market cap >= USD 250M; listed ADV > USD 5M over 6M; no OFAC-sanctioned countries; Bloomberg searchable.", f_note)

# ---------------- Term Sheet ----------------
wt = sheet("Term Sheet", "Indicative term sheet (links to Pricer)", [34, 70])
terms = [("Issuer", "Natixis (guaranteed by BPCE)"), ("Client", "NKE Private Wealth (Chak family office)"),
         ("Notional", "=TEXT(Pricer!B5,\"$#,##0\")"), ("Currency", "USD; underlying quanto (no FX exposure)"),
         ("Trade / issue date", "17 Sep 2026"), ("Maturity", "=\"Sep \"&(2026+Pricer!B6)&\" (\"&Pricer!B6&\"Y)\""),
         ("Underlying", "Natixis Asia Digital Bridge 10% VT Index (USD, 1% decrement), on an equal-weight basket of 13 Asian AI-hardware stocks"),
         ("Index rules", "Daily exposure = min(10% / max(20d, 60d realised vol), 150%), cash on the remainder, minus 1% p.a."),
         ("Capital protection", "=TEXT(Pricer!B7,\"0%\")&\" at maturity, subject to issuer credit\""),
         ("Participation", "=TEXT(Pricer!B23,\"0%\")&\" of index performance, uncapped\""),
         ("Redemption", "Protection + Participation x max(0, Index final / Index initial - 1)"),
         ("Optional enhancements", "12M final averaging (raises participation); Legacy Lock 130/160/190 (lowers it); see MC Results"),
         ("Appendix variant", "Best-of with a spot Bitcoin ETF sleeve (IBIT US); not in the core proposal"),
         ("Secondary market", "Natixis indicative daily bid, under normal market conditions"),
         ("Pricing basis", "Funding from the game-rules grid; option inputs are placeholders to be refreshed on Bloomberg")]
for i, (a, b) in enumerate(terms):
    put(wt, f"A{3 + i}", a, f_b); put(wt, f"B{3 + i}", b, f_link if str(b).startswith("=") else f_calc)
    wt[f"B{3 + i}"].alignment = Alignment(wrap_text=True)

from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)   # Excel computes every formula on open
wb.save(OUT)
print("saved", OUT)
