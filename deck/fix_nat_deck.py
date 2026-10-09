"""Fix the inconsistencies in the team deck (NAT.pptx) so every slide describes the same note:
annual step-down autocall, call level 105% - 2.5% x year (102.5% ... 92.5%), paid 100% + 12.75% x years at call,
70% barrier observed at Year 5 only, client coupon 12.75%, fair coupon 14.092219% (Appendix E).

Re-run figures (annual model, Appendix F inputs, 400k paths; coupon shifts measured on that model and applied
to the 14.09% base): see pricing/scenario2_annual_checks.py.

Usage: python3 deck/fix_nat_deck.py <in.pptx> <out.pptx>
"""
import sys
import numpy as np
from pptx import Presentation
from pptx.chart.data import XyChartData
from pptx.util import Inches

CPN, FAIR, SE_BP = 12.75, 14.092219, 2.236
FV = 1 - (FAIR - CPN) / 100 * 1.22            # PV of the note at 12.75% (E[b] = 1.22 discounted coupon-years)
LO, HI = FAIR - 1.96 * SE_BP / 100, FAIR + 1.96 * SE_BP / 100


# ---------------------------------------------------------------- helpers
def shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(name)


def set_tf(tf, paras):
    """Replace a text frame's text paragraph by paragraph, keeping each paragraph's first-run formatting."""
    if isinstance(paras, str):
        paras = [paras]
    ps = list(tf.paragraphs)
    while len(ps) < len(paras):  # clone the last paragraph
        import copy
        new = copy.deepcopy(ps[-1]._p); ps[-1]._p.addnext(new)
        ps = list(tf.paragraphs)
    for p, text in zip(ps, paras):
        runs = p.runs
        if not runs:
            p.add_run().text = text; continue
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    for p in ps[len(paras):]:
        p._p.getparent().remove(p._p)


def set_text(slide, name, paras):
    set_tf(shape(slide, name).text_frame, paras)


def set_cell(slide, name, r, c, text, table_index=0):
    tables = [sh for sh in slide.shapes if sh.name == name and sh.has_table]
    set_tf(tables[table_index].table.cell(r, c).text_frame, text)


def set_row(slide, name, r, values, table_index=0):
    for c, v in enumerate(values):
        set_cell(slide, name, r, c, v, table_index)


def xy(chart, series):
    d = XyChartData()
    for nm, pts in series:
        s = d.add_series(nm)
        for x, y in pts:
            s.add_data_point(x, y)
    chart.replace_data(d)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def drop_stray_page_numbers(slide):
    """The layout already numbers the slide; remove leftover numeric text boxes in the bottom-right corner."""
    for sh in list(slide.shapes):
        if sh.has_text_frame and sh.text_frame.text.strip().isdigit() and sh.left > Inches(12) and sh.top > Inches(6.8):
            sh._element.getparent().remove(sh._element)


# ---------------------------------------------------------------- fixes
def fix(prs):
    S = prs.slides

    # Slide 2: banner grammar
    for sh in S[1].shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith("The family owns Asian infrastructure"):
            set_tf(sh.text_frame, "The family already owns Asian infrastructure, renewables and property: we add the AI hardware layer, not more of what they own")

    for i in (5, 6, 7, 9):
        drop_stray_page_numbers(S[i])

    # Slide 6: term sheet
    s = S[5]
    set_text(s, "Text 7", "Why this design: a double-digit coupon, a 70% barrier at maturity, and capital usually back early")
    set_cell(s, "Table 0", 5, 1, f"{CPN}% p.a., paid at call: 100% + {CPN}% × years elapsed")
    set_cell(s, "Table 0", 6, 1, "Called in year k if basket ≥ 105% − 2.5% × k (102.5%, 100%, 97.5%, 95%, 92.5%)")
    set_cell(s, "Table 0", 7, 0, "Maturity barrier")
    set_cell(s, "Table 0", 7, 1, "70%, observed at Year 5 only: if never called and the basket ends below 70%, repaid at the basket level; otherwise 100%")
    set_cell(s, "Table 0", 8, 1, f"100% / {FV*100:.1f}% (fair coupon {FAIR:.2f}%)")
    notes(s, f"The note: five years, five Asian ETFs, checked once a year. If the basket is at or above that year's call level "
             f"(102.5% in year 1, stepping down 2.5% a year to 92.5%), the note ends and pays 100% plus {CPN}% for every year elapsed. "
             f"If it is never called, the family gets 100% back unless the basket ends below 70% at year 5, in which case it is repaid "
             f"at the basket level. The fair coupon is {FAIR:.2f}%; we offer {CPN}%, so the note is worth about {FV*100:.1f}% at issue.")

    # Slide 9: how it pays
    s = S[8]
    set_text(s, "Text 5", "How it pays: one question every year, three outcomes")
    set_text(s, "Text 7", "Is the basket at or above that year's call level? (checked once a year: 105% − 2.5% × year)")
    set_text(s, "Text 8", f"Call levels step down from 102.5% (Y1) to 92.5% (Y5). Called in year k: 100% + {CPN}% × k. No coupon before the call.")
    set_text(s, "Text 15", ["If ≥ 102.5%: ", "CALLED,", "Receive 112.75% total, ", "Otherwise, refer to B,C "])
    set_text(s, "Text 19", ["If ≥ 100% & Y1 not ", "called: CALLED,", "Receive 125.5% total", "Otherwise, refer to B,C"])
    set_text(s, "Text 23", ["If ≥ 97.5% & Y1–2 not called: CALLED,", "Receive 138.25% total", "Otherwise, refer to B,C"])
    set_text(s, "Text 27", ["If ≥ 95% & Y1–3 not ", "called: CALLED", "Receive 151.0% total; ", "Otherwise, refer to B,C"])
    set_text(s, "Text 31", ["If ≥ 92.5% & Y1–4 not called: CALLED,", "Receive 163.75% total; ", "Otherwise refer to B,C"])
    set_text(s, "Text 37", f"Basket ≥ that year's call level: 100% plus {CPN}% × years. The note ends.")
    set_text(s, "Text 38", ["99.3% called", "within 5 years"])
    set_text(s, "Text 42", "Not called, basket ≥ 70% at Y5")
    set_text(s, "Text 43", ["Years 1–4: continue holding", "Year 5: repaid at 100% of notional"])
    set_text(s, "Text 44", "0.5% not called but ≥ 70% at Y5")
    set_text(s, "Text 49", "Repaid at the basket level at Y5: the family bears the full fall.")
    set_text(s, "Text 50", "0.2% end below 70% at Y5")
    set_text(s, "文字方塊 1", "Expected value at exit (team outlook): ~114.2% of notional")
    notes(s, f"Explain the note as one question asked once a year: is the basket at or above this year's call level? The level starts at "
             f"102.5% and steps down 2.5% a year to 92.5%. If yes, the note ends and the family receives 100% plus {CPN}% for every year "
             f"elapsed. If the note is never called, the family gets 100% back unless the basket ends below 70% at year 5. "
             f"Probabilities are on the team outlook (expected returns from the strategy app), not the pricing assumptions.")

    # Slide 10: risks and hedging
    s = S[9]
    set_row(s, "Table 0", 1, ["Capital loss",
                              "Never called and basket < 70% at Y5: repaid at the basket level (0.2% on the team outlook; ~14% under pricing assumptions, average loss ~50%)",
                              "Diversified basket; 70% barrier at maturity only; ~5% of family wealth"])
    set_text(s, "Text 10", "Natixis owes digital coupons on the step-down call levels (102.5% → 92.5%): replicated with tight basket call spreads, delta-hedged with ETF units, constituent stocks and index futures as proxies; listed options on EWY/EWT")
    set_text(s, "Text 13", "Coupon value moves with basket vol and correlation (fair 12.6%–15.6% for σ ± 3pts): managed with listed index options and correlation trades")
    set_text(s, "Text 19", "Natixis is long the client's put: struck at 100%, live only if never called and the basket ends below 70% at Y5; hedged with basket put spreads and delta. Principal funded at 5.69%")
    set_text(s, "Text 20", ["Pricing sensitivity (fair coupon, annual model):",
                            f"Base case: C_fair = {FAIR:.2f}%",
                            "Volatility shift: σ ± 3pts ⟹ C_fair ∈ [12.6%, 15.6%], ∂C/∂σ > 0",
                            "Correlation shift: ρ −0.15 / +0.14 ⟹ C_fair ∈ [13.0%, 15.0%], ∂C/∂ρ > 0",
                            "Equity drift: Δμ = −1% to −2% ⟹ C_fair = 15.3%–16.7%",
                            f"The {CPN}% offer stays below fair value in every case except vol −3pts (12.6%)."])
    notes(s, f"Investor risks on the left, desk hedges on the right. The capital-loss line is explicit: the family loses only if the note is "
             f"never called and the basket ends below 70% at year 5. That is 0.2% of paths on our outlook but about 14% under the pricing "
             f"assumptions, which is what the coupon pays for. The {CPN}% coupon sits below the {FAIR:.2f}% fair level in every sensitivity "
             f"except a 3-point fall in volatility.")

    # Appendix A: assumptions (now the inputs actually used, from Appendix F)
    s = S[10]
    set_text(s, "Text 7", "Inputs from Appendix F (implied vols, correlations); the strategy app (slides 7–8) uses historical data and is not used for pricing")
    vol_rng = ["30–38%", "42–45%", "32–36%", "23–26%", "7–12%"]
    rho = [["1.00", "0.53", "0.74", "0.86", "0.56"], ["0.53", "1.00", "0.79", "0.69", "0.44"], ["0.74", "0.79", "1.00", "0.61", "0.69"],
           ["0.86", "0.69", "0.61", "1.00", "0.52"], ["0.56", "0.44", "0.69", "0.52", "1.00"]]
    set_cell(s, "Table 0", 0, 2, "Vol Y1–Y5"); set_cell(s, "Table 0", 0, 3, "Div")
    for i in range(5):
        set_cell(s, "Table 0", i + 1, 2, vol_rng[i]); set_cell(s, "Table 0", i + 1, 3, "0%")
        for j in range(5):
            set_cell(s, "Table 0", i + 1, 4 + j, rho[i][j])
    set_text(s, "Text 8", [
        "Model. Correlated GBM for each ETF, 1,000,000 paths, annual steps, implied-vol term structure through interval variances (App. F). Risk-neutral drift = USD funding; dividends set to zero. Quanto adjustment ignored (assumed hedged by the desk).",
        "Discounting. Case USD funding grid; 5Y discount factor 0.758.",
        f"Coupon solve. Fair coupon makes the note worth exactly 100% (App. E): {FAIR:.2f}%. We offer {CPN}%; the {FAIR - CPN:.2f}% a year difference is Natixis' hedging cost and margin.",
        "Historical back-test (to run). Monthly Bloomberg history since 2021 for all five ETFs. Launch a hypothetical note every month, apply the same call levels and barrier, and record coupons, life and IRR by launch year."])
    var = [("Variant of our note", "Fair cpn"), ("70% barrier at Y5 (base)", f"{FAIR:.1f}%"), ("Barrier 65%", "13.8%"), ("Barrier 75%", "14.3%"),
           ("No barrier", "14.7%"), ("Equity drift −1%", "15.3%"), ("Equity drift −2%", "16.7%"), ("Vol +3pts / −3pts", "15.6% / 12.6%")]
    for r, row in enumerate(var):
        set_row(s, "Table 1", r, list(row))
    set_text(s, "Text 9", f"Variants re-run on the annual model (400k paths) and shown as shifts from the {FAIR:.2f}% base. Code: pricing/scenario2_annual_checks.py")

    # Appendix B: scenarios on the annual note
    s = S[11]
    set_text(s, "Text 7", "Hypothetical basket paths, checked yearly against the call levels (102.5% → 92.5%) and the 70% barrier at Y5")
    rows = [("AI capex boom", "Called Y1 (118% ≥ 102.5%)", "12.75%", "112.75", "12.75%", "–"),
            ("Range-bound Asia", "Called Y2 (101% ≥ 100%)", "25.5%", "125.5", "12.0%", "–"),
            ("2022-style tech drawdown, slow recovery", "Called Y4 (104% ≥ 95%)", "51.0%", "151.0", "10.9%", "–"),
            ("Lost decade (sideways below par)", "Not called; 88% ≥ 70%: 100% back", "0.0%", "100.0", "0.0%", "−12%"),
            ("Geopolitical shock, no recovery", "Not called; 55% < 70%", "0.0%", "55.0", "−11.3%", "−45%")]
    for r, row in enumerate(rows, 1):
        set_row(s, "Table 0", r, list(row))
    ch = shape(s, "Chart 0").chart
    if ch.has_title:
        set_tf(ch.chart_title.text_frame, "Basket level (% of initial); call levels 102.5% → 92.5%, 70% barrier at Y5")
    set_text(s, "Text 9", "Reading: the note earns 10.9%–12.75% a year whenever the basket reaches that year's call level, even after a 30% drawdown (called in Y4). In the lost decade the basket ends at 88%, above the 70% barrier, so the family gets 100% back; only a lasting fall below 70% passes the loss through (55%).")
    notes(s, "Walk the five scenarios. The 2022-style path is the key one: a 30% drawdown in the first year that would scare a direct investor, yet the note is called in year 4 at 151% because the call level has stepped down to 95%. The lost decade ends at 88%, above the barrier, so the family gets its money back. Only a shock that leaves the basket below 70% at year 5 costs capital.")

    # Appendix C: replication
    s = S[12]
    shape(s, "Text 5").top = Inches(0.11)
    set_text(s, "Text 10", "100% repaid at call or maturity. Worth 88.2%")
    set_text(s, "Text 14", f"{CPN}% × years, paid on the call date when the basket ≥ that year's level. Worth 15.6%")
    set_text(s, "Text 18", "Each year, basket ≥ 105% − 2.5% × year: note ends, future coupons cancelled. Keeps the coupon affordable")
    set_text(s, "Text 22", "If never called and the basket ends below 70% at Y5, the client bears the full fall from 100%. Worth −5.4%")
    set_text(s, "Text 24", f"Fair value {FV*100:.1f}% = principal 88.2% + coupons 15.6% − short put 5.4%.  Issued at 100%, leaving {100 - FV*100:.1f}% "
                           f"(≈{FAIR - CPN:.2f}% a year) for hedging and margin. The short put funds the double-digit coupon.")
    xs = list(range(40, 161))
    xy(shape(s, "Chart 1").chart, [("Coupon digitals", [(x, CPN if x >= 102.5 else 0) for x in xs])])
    va = shape(s, "Chart 1").chart.value_axis; va.minimum_scale, va.maximum_scale = 0, 15
    xy(shape(s, "Chart 2").chart, [("Autocall (Natixis)", [(x, 1 if x < 102.5 else 0) for x in xs])])
    xy(shape(s, "Chart 3").chart, [("Short put (client)", [(x, x - 100 if x < 70 else 0) for x in xs])])
    notes(s, "Mini-charts show each leg against the basket level (year-1 view for the coupon and autocall legs). The short put is the key: the client sells Natixis a put struck at 100% that only applies if the note is never called and the basket ends below 70%, and its premium funds the double-digit coupon.")

    # Appendix D: alternatives
    s = S[13]
    shape(s, "Text 5").top = Inches(0.11)
    set_text(s, "Text 7", "Same basket and call levels: full protection, no barrier, and a pure participation note")
    xs = list(range(40, 181))
    old = shape(s, "Chart 0").chart
    ppn = [(x, 100 + max(0, x - 100) * 0.853) for x in xs]
    xy(old, [(f"Our note ({CPN}%)", [(x, x if x < 70 else (100 if x < 92.5 else 100 + 5 * CPN)) for x in xs]),
             ("100% protected (9.6%)", [(x, 100 if x < 92.5 else 148.0) for x in xs]),
             ("PPN", ppn), ("Basket", [(x, x) for x in xs])])
    rows = [("Option", "Return", "Loss risk"),
            ("Step-down autocall (ours)", f"{CPN}% p.a. at call", "14% pricing; 0.2% outlook"),
            ("100% protected autocallable", "~9.6% fair", "0%"),
            ("No barrier", "~14.7% fair", "~19% pricing"),
            ("Principal-protected note", "85% of gain at Y5", "0%")]
    for r, row in enumerate(rows):
        set_row(s, "Table 0", r, list(row))
    set_text(s, "Text 9", f"Trade-off: full protection pays only ~9.6%. Removing the barrier adds ~0.6% of coupon but loses capital in ~19% of paths instead of ~14% (pricing measure). The PPN pays nothing unless the basket ends above 100% after 5 years. Our note pays {CPN}%, with money usually back within ~1 year on the team outlook, and protects the family down to a 30% fall at maturity.")
    set_text(s, "Text 10", f"Annual Monte Carlo on App. F inputs (400k paths); fair coupons shown as shifts from the {FAIR:.2f}% base. Payoffs at Y5 for notes not called earlier. Inputs to refresh with Bloomberg data as of the 17 Sep 2026 trade date.")

    # Appendix E
    s = S[14]
    shape(s, "Text 0").top = Inches(0.11)
    set_cell(s, "Table 0", 1, 1, "105% − 2.5% × k: Y1 102.5% · Y2 100% · Y3 97.5% · Y4 95% · Y5 92.5%")
    set_cell(s, "Table 1", 2, 2, f"[{LO:.6f}%, {HI:.6f}%]")
    set_text(s, "Text 42", f"{FAIR:.6f}% fair = {CPN}% coupon to the client + {FAIR - CPN:.6f}% kept by Natixis")
    set_text(s, "Text 44", f"We keep {FAIR - CPN:.6f}% a year and give the client a {CPN}% coupon")

    # Appendix F: discount factors and interval variances recomputed from the tables' own inputs
    s = S[15]
    R = np.array([.0478, .053, .0552, .05605, .0569]); r = np.log1p(R); t = np.arange(1, 6)
    for i in range(5):
        set_cell(s, "Table 0", i + 1, 2, f"{r[i]*100:.3f}%"); set_cell(s, "Table 0", i + 1, 3, f"{np.exp(-r[i]*t[i]):.4f}", 0)
    vol = np.array([[.3045, .3234, .335, .367, .375], [.415, .427, .445, .445, .445], [.324, .358, .363, .363, .363],
                    [.23, .249, .256, .256, .259], [.07, .09, .12, .12, .12]])
    tv = vol ** 2 * t; dV = np.diff(np.c_[np.zeros(5), tv], axis=1)
    tbl = [sh for sh in s.shapes if sh.has_table]
    dv_table = [tb for tb in tbl if tb.table.cell(0, 0).text.startswith("Asset / Interval")][0].table
    for i in range(5):
        for j in range(5):
            set_tf(dv_table.cell(i + 1, j + 1).text_frame, f"{dV[i, j]:.4f}")
    set_text(s, "Text 17", "Basket weights 3119.HK 35% · EWY 20% · EWT 15% · 2644.T 15% · 00878.TW 15%   |   Call levels 105% − 2.5% × year (102.5% → 92.5%)   |   Maturity barrier 70%")
    set_text(s, "Text 18", "Source: team write-up, section 4. Rates: case USD funding grid (Year 4 interpolated); discount factors D = exp(−r·t) and ΔV = t·σ(t)² − (t−1)·σ(t−1)² recomputed from the tables. Correlation matrix checked positive definite.")


if __name__ == "__main__":
    prs = Presentation(sys.argv[1])
    fix(prs)
    prs.save(sys.argv[2])
    print("saved", sys.argv[2], "fair value at", CPN, "=", round(FV * 100, 2), "CI", round(LO, 6), round(HI, 6))
