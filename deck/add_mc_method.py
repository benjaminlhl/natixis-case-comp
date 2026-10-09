"""Add the Monte Carlo methodology box to the 'How it pays' slide of the team deck.
The payoff picture and call-probability chart are narrowed to make room for a box on the right.

Usage: python3 deck/add_mc_method.py <in.pptx> <out.pptx>
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR

BLUE, TXT, FILL = RGBColor(0x2C, 0x5F, 0x82), RGBColor(0x26, 0x26, 0x26), RGBColor(0xE8, 0xEF, 0xF5)
METHOD = [
    ("Historical resampling", "past 5 years of monthly returns, drawn at random with replacement."),
    ("Path generation", "10,000 five-year paths, each built from 60 monthly returns."),
    ("Payoff valuation", "each path runs the note's exact call levels, coupon and 70% barrier."),
    ("Risk & return metrics", "results aggregated into call probabilities, barrier breaches and the payout distribution."),
]


def find_slide(prs):
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.startswith("How it pays"):
                return s
    raise SystemExit("'How it pays' slide not found")


def main(src, dst):
    prs = Presentation(src)
    s = find_slide(prs)
    for sh in s.shapes:
        if sh.shape_type == 13:                                   # payoff picture
            sh.left, sh.top, sh.width, sh.height = Inches(7.05), Inches(3.2), Inches(3.5), Inches(3.5 * 2.47 / 4.54)
        elif sh.has_chart:                                         # cumulative call probability
            sh.left, sh.top, sh.width, sh.height = Inches(7.05), Inches(5.2), Inches(3.5), Inches(1.5)
        elif sh.has_text_frame and sh.text_frame.text.startswith("Expected value"):
            sh.width = Inches(6.7)

    x, y, w, h = 10.72, 3.15, 2.45, 3.55
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = FILL; box.line.fill.background(); box.shadow.inherit = False
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.06), Inches(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background(); bar.shadow.inherit = False

    tb = s.shapes.add_textbox(Inches(x + 0.16), Inches(y + 0.1), Inches(w - 0.26), Inches(h - 0.2))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    def run(p, text, bold=False, color=TXT, size=9):
        r = p.add_run(); r.text = text
        r.font.name = "Arial"; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color

    p = tf.paragraphs[0]; p.space_after = Pt(6)
    run(p, "Monte Carlo methodology", bold=True, color=BLUE, size=11)
    for head, body in METHOD:
        p = tf.add_paragraph(); p.space_after = Pt(5)
        run(p, "• " + head + ": ", bold=True); run(p, body)
    p = tf.add_paragraph()
    run(p, "Probabilities on this slide come from these paths (team outlook), not the pricing model in Appendix E.", color=RGBColor(0x59, 0x59, 0x59), size=8)
    prs.save(dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
